# Automated build source for the customer churn portfolio deck.
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

CSV = "customer_data.csv"
OUT = "customer_churn_analysis.pptx"
df = pd.read_csv(CSV)
assert len(df) == 30
assert df.isna().sum().sum() == 0
assert (df["Revenue_EGP"] == df["Tenure_Months"] * df["Monthly_Fee_EGP"]).all()
total_revenue = int(df["Revenue_EGP"].sum())
customers = len(df)
churned = int(df["Churned"].sum())
churn_rate = df["Churned"].mean()
plan = df.groupby("Plan").agg(customers=("Churned","size"), churned=("Churned","sum"), revenue=("Revenue_EGP","sum"))
plan["churn_rate"] = plan["churned"] / plan["customers"]
city = df.groupby("City").agg(customers=("Churned","size"), churned=("Churned","sum"))
city["churn_rate"] = city["churned"] / city["customers"]
channel = df.groupby("Channel").agg(customers=("Churned","size"), churned=("Churned","sum"))
channel["churn_rate"] = channel["churned"] / channel["customers"]
assert total_revenue == 316230
assert churned == 10
assert round(churn_rate, 4) == 0.3333
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
def slide(title, bullets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(0.7), Inches(0.55), Inches(12), Inches(0.8))
    p = box.text_frame.paragraphs[0]; p.text = title; p.font.size = Pt(28); p.font.bold = True
    y = 1.55
    for b in bullets:
        tb = s.shapes.add_textbox(Inches(0.9), Inches(y), Inches(11.5), Inches(0.55))
        p = tb.text_frame.paragraphs[0]; p.text = b; p.font.size = Pt(18); y += 0.72
slide("Customer Churn Analysis", ["Beginner portfolio project","Business question: Which customer segments show the highest observed churn risk?","Synthetic learning dataset • Jan–Mar 2026"])
slide("Dataset & Scope", [f"Customers: {customers}","Cities: Cairo, Giza, Qalyubia","Plans: Basic, Standard, Premium","Channels: Mobile and Web","Descriptive practice data — not a predictive model."])
slide("KPI Summary", [f"Revenue represented: {total_revenue:,} EGP",f"Customers: {customers}",f"Churned customers: {churned}",f"Overall churn rate: {churn_rate:.2%}"])
slide("Churn by Plan", [f"{idx}: {row.churn_rate:.2%} churn ({int(row.churned)}/{int(row.customers)} customers)" for idx,row in plan.sort_values("churn_rate",ascending=False).iterrows()])
slide("Churn by City", [f"{idx}: {row.churn_rate:.2%} churn ({int(row.churned)}/{int(row.customers)} customers)" for idx,row in city.sort_values("churn_rate",ascending=False).iterrows()])
slide("Churn by Channel", [f"{idx}: {row.churn_rate:.2%} churn ({int(row.churned)}/{int(row.customers)} customers)" for idx,row in channel.sort_values("churn_rate",ascending=False).iterrows()])
slide("Revenue by Plan", [f"{idx}: {int(row.revenue):,} EGP" for idx,row in plan.sort_values("revenue",ascending=False).iterrows()])
slide("Key Observations & Limitations", ["Basic has the highest observed churn rate.","Qalyubia has the highest observed city churn rate.","Mobile has higher observed churn than Web.","Small synthetic dataset: findings are descriptive and should not be treated as predictive."])
slide("Conclusion & Next Steps", ["Use the dashboard to explore segments.","Practice retention-focused SQL and Power BI analysis.","For real analysis, add more customers, tenure history, support interactions and cohort data."])
prs.save(OUT)
print(f"Created {OUT}")