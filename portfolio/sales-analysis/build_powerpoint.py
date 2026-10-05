from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

OUT = Path(__file__).with_name("sales_analysis_portfolio.pptx")
slides = [
("Sales Analysis","Beginner Portfolio Project\nPractice dataset • Jan–Mar 2026"),
("Dataset Overview","25 practice sales records\n7 columns: Date, City, Category, Product, Units, Unit Price, Revenue\nSynthetic learning dataset — not real company data"),
("KPI Summary","Total Revenue: 80,270 EGP\nTotal Units: 247\nOrders: 25\nAverage Order Value: 3,210.80 EGP (3,211 rounded)"),
("Revenue by Month","January 2026: 26,520 EGP\nFebruary 2026: 27,520 EGP\nMarch 2026: 26,230 EGP"),
("Revenue by Category","Electronics: 47,710 EGP\nHome: 17,200 EGP\nOffice: 15,360 EGP"),
("Revenue by City","Cairo: 32,910 EGP\nGiza: 28,110 EGP\nQalyubia: 19,250 EGP"),
("Top Products","Wireless Mouse: 14,400 EGP\nKeyboard: 12,600 EGP\nHeadphones: 8,550 EGP\nWebcam: 8,400 EGP\nDesk Lamp: 7,680 EGP"),
("Key Observations","Electronics is the highest-revenue category.\nFebruary is the strongest month by revenue.\nCairo generates the highest city revenue.\nWireless Mouse is the top-revenue product."),
("Conclusion","Clean/check data → calculate KPIs → aggregate → communicate findings.\nTools: Python/Pandas • SQL • Power BI-ready design")]
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
for i,(title,body) in enumerate(slides):
 s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor(248,250,252)
 t=s.shapes.add_textbox(Inches(.8),Inches(.7),Inches(11.8),Inches(1)).text_frame.paragraphs[0]; t.text=title; t.font.size=Pt(30); t.font.bold=True; t.font.color.rgb=RGBColor(15,23,42)
 tf=s.shapes.add_textbox(Inches(1),Inches(2),Inches(11),Inches(4.5)).text_frame; tf.word_wrap=True
 for j,line in enumerate(body.split("\n")):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph(); p.text=line; p.font.size=Pt(21); p.font.color.rgb=RGBColor(15,118,110); p.space_after=Pt(14)
 f=s.shapes.add_textbox(Inches(.8),Inches(6.9),Inches(11.8),Inches(.3)).text_frame.paragraphs[0]; f.text=f"Sales Analysis • {i+1}/{len(slides)}"; f.font.size=Pt(10); f.alignment=PP_ALIGN.RIGHT; f.font.color.rgb=RGBColor(100,116,139)
prs.save(OUT)
print(OUT)